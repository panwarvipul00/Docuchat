"""File -> list of pages of text.

Each parser returns `[(page_number, text), ...]` with 1-based page numbers, which
is what lets answers cite "document X, page N" later. Spreadsheets use one
"page" per sheet.
"""

import csv
import datetime as dt
import re
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, replace
from pathlib import Path

import openpyxl
import pymupdf
from django.conf import settings
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class Page:
    number: int
    text: str
    # Spreadsheet sheets: `text` is one record (row) per paragraph, and the
    # chunker keeps records whole and prefixes every chunk with `heading`
    # (e.g. 'Sheet "March"') so each chunk says where it came from.
    records: bool = False
    heading: str = ""


class ParsingError(Exception):
    """Raised for files we can't extract usable text from."""


def parse_file(path: str | Path, file_type: str) -> list[Page]:
    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "txt": parse_txt,
        "xlsx": parse_xlsx,
        "csv": parse_csv,
    }
    if file_type not in parsers:
        raise ParsingError(f"Unsupported file type: {file_type}")
    pages = clean_pages(parsers[file_type](Path(path)))
    if not any(p.text for p in pages):
        if file_type in ("xlsx", "csv"):
            raise ParsingError("The spreadsheet has no data rows.")
        raise ParsingError(
            "No extractable text found. Scanned/image-only PDFs need OCR, which isn't supported."
        )
    return pages


def parse_pdf(path: Path) -> list[Page]:
    try:
        with pymupdf.open(path) as pdf:
            if pdf.needs_pass:
                raise ParsingError("PDF is password-protected.")
            # sort=True orders text blocks top-to-bottom, left-to-right, which
            # fixes the jumbled reading order some PDFs produce.
            return [Page(i + 1, page.get_text("text", sort=True)) for i, page in enumerate(pdf)]
    except ParsingError:
        raise
    except Exception as exc:
        raise ParsingError(f"Could not read PDF: {exc}") from exc


def parse_docx(path: Path) -> list[Page]:
    """DOCX has no fixed pages (layout happens in Word), so we approximate them by
    splitting on explicit page breaks and the `lastRenderedPageBreak` markers Word
    saves from its last layout. Documents without those come out as one page."""
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise ParsingError(f"Could not read DOCX: {exc}") from exc

    pages: list[list[str]] = [[]]
    body = doc.element.body
    for block in body.iterchildren():
        if block.tag == qn("w:p"):
            parts: list[str] = []
            for node in block.iter():
                is_break = (node.tag == qn("w:br") and node.get(qn("w:type")) == "page") or (
                    node.tag == qn("w:lastRenderedPageBreak")
                )
                if is_break:
                    pages[-1].append("".join(parts))
                    parts = []
                    pages.append([])
                elif node.tag == qn("w:t") and node.text:
                    parts.append(node.text)
                elif node.tag == qn("w:tab"):
                    parts.append("\t")
            pages[-1].append("".join(parts))
        elif block.tag == qn("w:tbl"):
            # Flatten tables row by row so their text is still searchable.
            for row in block.iter(qn("w:tr")):
                cells = ["".join(t.text or "" for t in cell.iter(qn("w:t"))) for cell in row]
                pages[-1].append(" | ".join(c for c in cells if c))

    texts = ["\n".join(lines) for lines in pages]
    # Word often emits a rendered break right at the start; drop empty pages
    # at the edges but keep numbering of the rest contiguous.
    while len(texts) > 1 and not texts[0].strip():
        texts.pop(0)
    while len(texts) > 1 and not texts[-1].strip():
        texts.pop()
    return [Page(i + 1, t) for i, t in enumerate(texts)]


def parse_txt(path: Path) -> list[Page]:
    raw = path.read_bytes().decode("utf-8", errors="replace")
    # Form feed (\f) is the conventional plain-text page separator.
    chunks = raw.split("\f") if "\f" in raw else [raw]
    return [Page(i + 1, t) for i, t in enumerate(chunks)]


# ---------- Spreadsheets ----------
#
# RAG retrieves a handful of chunks, so a chunk must make sense on its own. A
# bare row like "Rahul | Sales | 50000" doesn't; "Name: Rahul | Dept: Sales |
# Salary: 50000" does, and embeds close to questions like "Rahul's salary".
# So each row is written as header: value pairs, one row per paragraph.
#
# Limitation (worth knowing in interviews): this suits lookup questions. Totals
# or averages over many rows need the whole sheet, which retrieval never sends.


def _format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()  # Excel stores plain dates as midnight
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 50000.0 -> "50000"
    # Flatten line breaks inside a cell so one row stays one record.
    return " ".join(str(value).split())


def rows_to_text(rows: Iterable[Sequence[object]], row_budget: list[int]) -> str:
    """Rows -> "Header: value | ..." records separated by blank lines.

    The first non-empty row is the header. Blank header cells fall back to the
    column letter ("Column C"). Empty rows are skipped. `row_budget` is a
    one-item list shared across sheets, so the row cap covers the whole file.
    """
    header: list[str] | None = None
    records: list[str] = []
    for raw in rows:
        cells = [_format_cell(v) for v in raw]
        if not any(cells):
            continue
        if header is None:
            header = [c or f"Column {get_column_letter(i + 1)}" for i, c in enumerate(cells)]
            continue
        row_budget[0] -= 1
        if row_budget[0] < 0:
            raise ParsingError(
                f"The spreadsheet has more than {settings.MAX_SPREADSHEET_ROWS:,} rows. "
                "Split it into smaller files."
            )
        pairs = [
            f"{header[i] if i < len(header) else f'Column {get_column_letter(i + 1)}'}: {c}"
            for i, c in enumerate(cells)
            if c
        ]
        records.append(" | ".join(pairs))
    return "\n\n".join(records)


def parse_xlsx(path: Path) -> list[Page]:
    """One page per worksheet (chart sheets have no cells and are skipped).

    read_only streams rows instead of loading the whole workbook, and together
    with the row cap keeps a small file that unzips to millions of cells from
    exhausting memory. data_only returns the value Excel last calculated for a
    formula rather than the formula text; files saved by tools that don't store
    those values come out with formula cells empty.
    """
    budget = [settings.MAX_SPREADSHEET_ROWS]
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ParsingError(
            f"Could not read XLSX (it may be password-protected or corrupt): {exc}"
        ) from exc
    try:
        pages: list[Page] = []
        for number, sheet in enumerate(wb.worksheets, start=1):
            text = rows_to_text(sheet.iter_rows(values_only=True), budget)
            pages.append(Page(number, text, records=True, heading=f'Sheet "{sheet.title}"'))
        return pages
    except ParsingError:
        raise
    except Exception as exc:
        raise ParsingError(f"Could not read XLSX: {exc}") from exc
    finally:
        wb.close()


def parse_csv(path: Path) -> list[Page]:
    # utf-8-sig drops the byte-order mark Excel adds when it saves "CSV UTF-8".
    raw = path.read_bytes().decode("utf-8-sig", errors="replace")
    try:
        # Excel in many locales (e.g. most of Europe) saves with ";" or tabs.
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            raw[:8192], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(raw.splitlines(), dialect)
    text = rows_to_text(rows, [settings.MAX_SPREADSHEET_ROWS])
    return [Page(1, text, records=True)]


# ---------- Cleaning ----------

_WS_RE = re.compile(r"[ \t ]+")
_DIGITS_RE = re.compile(r"\d+")
_HYPHEN_BREAK_RE = re.compile(r"(\w)-\n(\w)")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _HYPHEN_BREAK_RE.sub(r"\1\2", text)  # re-join words hyphenated across lines
    lines = [_WS_RE.sub(" ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _line_key(line: str) -> str:
    # "Page 3 of 10" and "Page 4 of 10" should count as the same footer.
    return _DIGITS_RE.sub("#", line.strip().lower())


def strip_repeated_headers_footers(pages: list[Page], edge_lines: int = 2) -> list[Page]:
    """Remove lines that appear in the first/last `edge_lines` lines of most pages.

    Running headers/footers ("ACME Corp — Confidential", "Page 3 of 10") add
    noise to every chunk and pull unrelated chunks together in vector space.
    Only applied with 3+ pages, where "repeated" is meaningful.
    """
    if len(pages) < 3:
        return pages

    counts: Counter[str] = Counter()
    for page in pages:
        lines = [ln for ln in page.text.split("\n") if ln.strip()]
        edges = set(lines[:edge_lines] + lines[-edge_lines:])
        counts.update({_line_key(ln) for ln in edges})

    threshold = max(3, int(len(pages) * 0.6))
    repeated = {key for key, n in counts.items() if n >= threshold}
    if not repeated:
        return pages

    cleaned: list[Page] = []
    for page in pages:
        lines = page.text.split("\n")
        non_empty = [i for i, ln in enumerate(lines) if ln.strip()]
        edge_idx = set(non_empty[:edge_lines] + non_empty[-edge_lines:])
        kept = [
            ln for i, ln in enumerate(lines) if not (i in edge_idx and _line_key(ln) in repeated)
        ]
        cleaned.append(Page(page.number, "\n".join(kept).strip()))
    return cleaned


def clean_pages(pages: list[Page]) -> list[Page]:
    cleaned = [replace(p, text=normalize_text(p.text)) for p in pages]
    # Header/footer stripping is for printed pages; on spreadsheets it could
    # delete a legitimately repeated row from several sheets.
    if any(p.records for p in cleaned):
        return cleaned
    return strip_repeated_headers_footers(cleaned)
